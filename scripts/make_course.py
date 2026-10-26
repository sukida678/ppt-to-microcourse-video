#!/usr/bin/env python3
"""make_course — turn a deck into a narrated micro-lecture video.

Two phases, so the *agent* can write the teaching script in between:

  1. extract   deck -> <work>/png/*.png + <work>/slides.json (text + speaker notes)
  2. build     <work>/narration.json + png -> TTS -> <work>/course.mp4

`all` runs both, using an existing narration.json if there is one.

Inputs can be .pptx / .ppt / .pdf / .odp (rendered via LibreOffice) or a directory
of pre-rendered .png/.jpg slides (no office suite needed).

Examples
    python make_course.py extract --deck lesson.pptx --work build
    # ... agent writes build/narration.json ...
    python make_course.py build  --work build
    python make_course.py all    --deck lesson.pptx --work build --voice zh-CN-YunxiNeural

Only python-pptx (text), pymupdf (pdf->png), ffmpeg (video) and edge-tts (speech)
are used -- all cross-platform and free. No vendor/agent-platform APIs.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from urllib.parse import quote

HERE = Path(__file__).resolve().parent
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp"}


def safe_stdio() -> None:
    """Never let a `print()` abort the run on a legacy Windows code page.

    cp936 can encode Chinese, but not everything that reaches stdout: a work dir with
    a non-ASCII path plus a child process' captured output is enough to produce a
    character the code page rejects (a GBK-encoded '微' decoded as UTF-8 becomes
    U+03A2). That used to raise UnicodeEncodeError *after* the TTS had succeeded and
    kill the build. Keep the console's own encoding, just stop it from throwing.

    When stdout is a pipe it is being read by an agent, not a console: write UTF-8 so
    the JSON we print (paths included) decodes cleanly on the other side.
    """
    opts = {"errors": "replace"}
    try:
        if not sys.stdout.isatty():
            opts["encoding"] = "utf-8"
    except Exception:
        opts["encoding"] = "utf-8"
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(**opts)
        except Exception:  # a replaced stdout (pipe/logger) may not support it
            pass


def child_env() -> dict:
    """Give python children a UTF-8 stdout so we can decode what they print.

    `subprocess(text=True, encoding="utf-8")` is only correct if the child actually
    writes UTF-8; by default a piped child writes the ANSI code page and a Chinese
    path comes back as mojibake.
    """
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    return env


safe_stdio()

SOFFICE_CANDIDATES = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    r"%LOCALAPPDATA%\Programs\LibreOffice\program\soffice.exe",
    r"D:\LibreOffice\program\soffice.exe",
    r"C:\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/bin/libreoffice",
    "/usr/local/bin/soffice",
    "/opt/libreoffice/program/soffice",
    "/snap/bin/libreoffice",
]


def find_soffice(explicit: str | None = None) -> str | None:
    for cand in [explicit, os.environ.get("SOFFICE"),
                 os.environ.get("SOFFICE_BIN"),
                 os.environ.get("LIBREOFFICE_BIN")]:
        if cand and Path(os.path.expandvars(cand)).exists():
            return str(Path(os.path.expandvars(cand)))
    for name in ("soffice", "libreoffice", "soffice.exe"):
        hit = shutil.which(name)
        if hit:
            return hit
    for cand in SOFFICE_CANDIDATES:
        p = Path(os.path.expandvars(cand))
        if p.exists():
            return str(p)
    # last resort: portable copies next to the working dir or in /opt/libreoffice*
    for pattern in ("/opt/libreoffice*/program/soffice", "**/LibreOffice*/program/soffice.exe"):
        try:
            for p in Path.cwd().glob(pattern):
                return str(p)
        except Exception:
            pass
    return None


def run_cmd(cmd, cwd=None, timeout=900, label=""):
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8",
                         errors="replace", cwd=cwd, timeout=timeout, env=child_env())
    if res.returncode != 0:
        tail = "\n".join(((res.stderr or "") + (res.stdout or "")).strip().splitlines()[-12:])
        raise RuntimeError(f"{label or cmd[0]} failed (exit {res.returncode}):\n{tail}")
    return res


# --------------------------------------------------------------------------- extract
def _ppt2png():
    """The converter lives in ppt2png.py — one implementation, not two.

    make_course used to carry its own copy of the soffice call and the PyMuPDF
    renderer; the two copies drifted (ppt2png passed no profile, so a deck path with
    spaces broke there). Both now go through the functions below.
    """
    scripts_dir = Path(__file__).resolve().parent
    if str(scripts_dir) not in sys.path:
        sys.path.insert(0, str(scripts_dir))
    import ppt2png  # noqa: PLC0415

    return ppt2png


def deck_to_pdf(deck: Path, out_dir: Path, soffice: str) -> Path:
    """PPTX/PPT -> PDF with an isolated, percent-encoded LibreOffice profile.

    A literal space in the profile path used to kill the conversion with an empty
    message; `ppt2png.profile_uri` encodes it, and `soffice_hint` explains the
    failure when the converter still exits without printing anything.
    """
    pdf, _backend = _ppt2png().convert_to_pdf(
        deck, out_dir, "auto", soffice, profile=out_dir / "_lo_profile")
    return pdf


def pdf_to_png(pdf: Path, png_dir: Path, dpi: int) -> list[Path]:
    """PDF -> slide PNGs through the one renderer (drops stale PNGs first)."""
    png_dir.mkdir(parents=True, exist_ok=True)
    for old in png_dir.glob("*.png"):
        old.unlink()
    return _ppt2png().render_pdf(pdf, png_dir, dpi, "slide")


def extract_text(deck: Path) -> list[dict]:
    """Best-effort per-slide text + speaker notes from a .pptx."""
    try:
        from pptx import Presentation  # noqa: PLC0415
    except Exception:
        return []
    try:
        prs = Presentation(str(deck))
    except Exception:
        return []
    out = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        bullets: list[str] = []
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
            if not title and (is_title or (len(text) <= 60 and "\n" not in text)):
                title = text
            else:
                bullets.extend(ln.strip() for ln in text.splitlines() if ln.strip())
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        out.append({
            "index": idx,
            "title": title,
            "bullets": bullets,
            "notes": notes,
            "text": "\n".join([t for t in [title, *bullets] if t]),
        })
    return out


def cmd_extract(args) -> int:
    work = Path(args.work).resolve()
    work.mkdir(parents=True, exist_ok=True)
    png_dir = work / "png"
    src = Path(args.deck).resolve() if args.deck else None
    if not src or not src.exists():
        raise RuntimeError(f"deck not found: {src}")

    if src.is_dir():
        files = sorted(p for p in src.iterdir() if p.suffix.lower() in IMG_EXT)
        if not files:
            raise RuntimeError(f"no images in {src}")
        png_dir.mkdir(parents=True, exist_ok=True)
        for old in png_dir.glob("*.png"):
            old.unlink()
        images = []
        for i, f in enumerate(files, start=1):
            dst = png_dir / f"slide_{i:02d}.png"
            shutil.copy2(f, dst)
            images.append(dst)
        slides = [{"index": i + 1, "title": "", "bullets": [], "notes": "", "text": ""}
                  for i in range(len(images))]
        renderer = "images"
    else:
        slides = extract_text(src) if src.suffix.lower() == ".pptx" else []
        soffice = find_soffice(args.soffice)
        if not soffice:
            raise RuntimeError(
                "LibreOffice not found. Install it (brew/apt/winget install libreoffice), "
                "or pass --soffice /path/to/soffice, or export the slides to PNG yourself "
                "and pass --deck <dir-of-png>.")
        pdf = deck_to_pdf(src, work, soffice)
        images = pdf_to_png(pdf, png_dir, args.dpi)
        renderer = f"libreoffice:{Path(soffice).name}"

    if len(slides) != len(images):
        while len(slides) < len(images):
            slides.append({"index": len(slides) + 1, "title": "", "bullets": [],
                           "notes": "", "text": ""})
        slides = slides[: len(images)]

    for s, img in zip(slides, images):
        s["image"] = str(img.relative_to(work).as_posix())

    (work / "slides.json").write_text(
        json.dumps({"renderer": renderer, "dpi": args.dpi, "slides": slides},
                   ensure_ascii=False, indent=1), encoding="utf-8")

    _write_narration_template(work, slides)
    print(json.dumps({
        "ok": True,
        "renderer": renderer,
        "slides": len(slides),
        "images": str(png_dir),
        "slides_json": str(work / "slides.json"),
        "narration_template": str(work / "narration.json"),
        "next": "edit narration.json (one line of spoken text per slide), then run: "
                f"build --work {work}",
    }, ensure_ascii=False, indent=1))
    return 0


def _write_narration_template(work: Path, slides: list[dict]) -> None:
    target = work / "narration.json"
    if target.exists():
        return
    body = {}
    for s in slides:
        hint = s.get("notes") or " ".join(s.get("bullets") or []) or s.get("title") or ""
        body[str(s["index"])] = hint
    target.write_text(json.dumps({
        "voice": "zh-CN-XiaoxiaoNeural",
        "_comment": "把每页换成口语化的讲解词（60-120 字，一页约 15-30 秒）。空字符串=该页无旁白。",
        "slides": body,
    }, ensure_ascii=False, indent=1), encoding="utf-8")


# ----------------------------------------------------------------------------- build
def load_narration(work: Path) -> tuple[str | None, dict[str, str]]:
    p = work / "narration.json"
    if not p.exists():
        return None, {}
    data = json.loads(p.read_text(encoding="utf-8-sig"))
    voice = data.get("voice") if isinstance(data, dict) else None
    raw = data.get("slides", data) if isinstance(data, dict) else data
    out: dict[str, str] = {}
    if isinstance(raw, dict):
        for k, v in raw.items():
            if str(k).startswith("_"):
                continue
            out[str(k)] = (v or "").strip() if isinstance(v, str) else ""
    elif isinstance(raw, list):
        for i, v in enumerate(raw, start=1):
            if isinstance(v, dict):
                k = str(v.get("index") or i)
                out[k] = (v.get("narration") or v.get("text") or "").strip()
            else:
                out[str(i)] = (v or "").strip()
    return voice, out


def _last_json(text: str) -> dict | None:
    """The compositor prints its summary as the *last* JSON object on stdout."""
    lines = text.strip().splitlines()
    for i in range(len(lines) - 1, -1, -1):
        if lines[i].lstrip().startswith("{"):
            try:
                return json.loads("\n".join(lines[i:]))
            except json.JSONDecodeError:
                return None
    return None


def _missing_manifest_refs(segments: list[dict], work: Path) -> list[tuple[str, str]]:
    """(id, path) for every file the manifest names but the disk does not have."""
    out = []
    for e in segments:
        for key in ("audio", "srt"):
            rel = e.get(key)
            if rel and not (work / rel).exists():
                out.append((str(e.get("id")), rel))
    return out


def cmd_build(args) -> int:
    work = Path(args.work).resolve()
    sj = work / "slides.json"
    if not sj.exists():
        raise RuntimeError(f"run extract first: {sj} missing")
    data = json.loads(sj.read_text(encoding="utf-8-sig"))
    slides = data["slides"]

    narr_path = work / "narration.json"
    if not narr_path.exists() and not args.allow_silent:
        raise RuntimeError(
            f"narration.json not found: {narr_path}\n"
            "  extract wrote a template there — replace each slide with spoken narration.\n"
            "  (pass --allow-silent to build a silent slideshow instead)")

    voice_cfg, narr = load_narration(work)
    voice = args.voice or voice_cfg or "zh-CN-XiaoxiaoNeural"

    source_text = {str(s["index"]):
                   (s.get("notes") or " ".join(s.get("bullets") or []) or s.get("title") or "").strip()
                   for s in slides}
    spoken = [v for v in narr.values() if v]
    if spoken and sum(1 for k, v in narr.items() if v and v == source_text.get(k)) >= 0.7 * len(spoken):
        print("  ! narration.json still matches the slide text verbatim — that is a screen "
              "reader, not a lesson. Rewrite it in spoken language.", flush=True)

    build_dir = work / "build"
    build_dir.mkdir(parents=True, exist_ok=True)
    segments = []
    empty = []
    for s in slides:
        sid = f"{s['index']:02d}"
        text = narr.get(str(s["index"]), "") or ""
        if not text:
            empty.append(s["index"])
        segments.append({"id": sid, "image": s["image"], "text": text,
                         "title": s.get("title", "")})
    seg_path = build_dir / "segments.json"
    seg_path.write_text(json.dumps(segments, ensure_ascii=False, indent=1), encoding="utf-8")

    if args.dry_run:
        print(json.dumps({"ok": True, "dry_run": True, "segments": len(segments),
                          "without_narration": empty, "segments_json": str(seg_path)},
                         ensure_ascii=False, indent=1))
        return 0

    # 1) speech
    tts_script = HERE / "tts_edge.py"
    tts_out = build_dir
    tts_cmd = [sys.executable, str(tts_script), "--segments", str(seg_path),
               "--out", str(tts_out), f"--voice={voice}",
               f"--rate={args.rate}", f"--pitch={args.pitch}"]
    if args.strict:
        tts_cmd.append("--strict")
    tts_run = run_cmd(tts_cmd, label="edge-tts", timeout=3600)
    print((tts_run.stdout or "").strip(), flush=True)
    tts = {str(r["id"]): r for r in json.loads(
        (tts_out / "tts.json").read_text(encoding="utf-8"))}
    failed = [r for r in tts.values() if r.get("error")]
    silent = [str(r["id"]) for r in tts.values() if r.get("silent")]

    # 2) manifest for the compositor
    man_segments = []
    for seg in segments:
        t = tts.get(seg["id"], {})
        entry = {"id": seg["id"], "image": seg["image"],
                 "duration": float(t.get("duration") or 0.0)}
        if t.get("audio"):
            entry["audio"] = str(Path(t["audio"]).resolve().relative_to(work).as_posix())
        if t.get("srt"):
            entry["srt"] = str(Path(t["srt"]).resolve().relative_to(work).as_posix())
        man_segments.append(entry)
    manifest = {"out": args.out or "course.mp4", "segments": man_segments}
    man_path = work / "manifest.json"
    man_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    missing_refs = _missing_manifest_refs(man_segments, work)
    if missing_refs:
        detail = ", ".join(f"[{sid}] {p}" for sid, p in missing_refs)
        print(f"  ! manifest points at {len(missing_refs)} missing file(s): {detail} "
              f"— the render will not be clean", flush=True)

    # 3) video
    compose = HERE / "compose.py"
    cmd = [sys.executable, str(compose), "--project", str(work),
           "--resolution", args.resolution, "--fps", str(args.fps),
           "--motion", args.motion, "--transition", args.transition,
           "--subs", args.subs,
           "--sub-size", str(args.sub_size), "--zoom", str(args.zoom)]
    if args.sub_font:
        cmd += ["--sub-font", args.sub_font]
    # --hold only means something if the compositor hears about it: it is the
    # duration of a slide without audio, and silently ignoring it made the log lie
    if "--default-hold" in compose.read_text(encoding="utf-8", errors="replace"):
        cmd += ["--default-hold", str(args.hold)]
    else:
        print("  ! this compose.py has no --default-hold — --hold is ignored, silent "
              "slides hold for the compositor's own default", flush=True)
    if args.sub_box:
        cmd.append("--sub-box")
    if args.bgm:
        cmd += ["--bgm", str(Path(args.bgm).resolve()), "--music-volume", str(args.music_volume)]
    out = run_cmd(cmd, label="compose", timeout=7200)
    lines = out.stdout.strip().splitlines()
    cut = next((i for i in range(len(lines) - 1, -1, -1)
                if lines[i].lstrip().startswith("{")), len(lines))
    print("\n".join(lines[:cut]).strip(), flush=True)
    summary = _last_json(out.stdout) or {}
    if failed:
        # the video still renders (those pages hold for --hold seconds) but the build
        # is not clean, and the JSON has to say so — that is what "ok" means
        print(f"  ! {len(failed)} page(s) produced no audio: "
              f"{', '.join(str(r['id']) for r in failed)} — they are silent in the video",
              flush=True)
        summary["tts_failed"] = [{"id": str(r["id"]), "error": r.get("error")}
                                 for r in failed]
        summary["ok"] = False
    if silent:
        summary["silent_slides"] = silent
    print(json.dumps(summary, ensure_ascii=False, indent=1) if summary else out.stdout.strip())
    if empty:
        print(f"  ! slides without narration (silent, {args.hold:g}s hold each): {empty}")
    return 0


def cmd_all(args) -> int:
    rc = cmd_extract(args)
    if rc != 0:
        return rc
    return cmd_build(args)


VALUE_OPTS = {"--deck", "--work", "--soffice", "--dpi", "--voice", "--rate", "--pitch",
              "--resolution", "--fps", "--motion", "--transition", "--subs", "--sub-font",
              "--sub-size", "--bgm", "--music-volume", "--hold", "--out"}


def normalize_argv(argv):
    """`--rate -6%` -> `--rate=-6%`.

    argparse reads a value that starts with '-' as another option, so the natural way
    to ask for a slower voice ("--rate -6%") fails with "expected one argument".
    Folding the pair into a single token makes both spellings work.
    """
    out, i = [], 0
    while i < len(argv):
        a = argv[i]
        if a in VALUE_OPTS and i + 1 < len(argv):
            out.append(f"{a}={argv[i + 1]}")
            i += 2
        else:
            out.append(a)
            i += 1
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="PPT/PDF/images -> narrated micro-lecture video",
        formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    ap.add_argument("phase", choices=["extract", "build", "all"])
    ap.add_argument("--deck", help="input .pptx/.ppt/.pdf/.odp, or a dir of slide images")
    ap.add_argument("--work", default="build", help="work dir (default: build)")
    ap.add_argument("--soffice", default=None, help="explicit path to the soffice binary")
    ap.add_argument("--dpi", type=int, default=150, help="slide raster dpi (default 150)")
    ap.add_argument("--voice", default=None, help="edge-tts voice (default from narration.json)")
    ap.add_argument("--rate", default="+0%")
    ap.add_argument("--pitch", default="+0Hz")
    ap.add_argument("--resolution", default="1920x1080")
    ap.add_argument("--fps", type=int, default=30)
    ap.add_argument("--motion", default="kenburns", choices=["kenburns", "none"])
    ap.add_argument("--transition", default="fade", choices=["fade", "none"])
    ap.add_argument("--subs", default="burn", choices=["burn", "soft", "none"])
    ap.add_argument("--sub-font", default=None,
                    help="subtitle font family (default: this platform's CJK font; "
                         "compose.py warns and substitutes when it is not installed)")
    ap.add_argument("--sub-size", type=int, default=21)
    ap.add_argument("--sub-box", action="store_true",
                    help="translucent band behind subtitles (slides with bottom-heavy content)")
    ap.add_argument("--zoom", type=float, default=1.06,
                    help="kenburns end zoom; use --motion none for full-bleed decks")
    ap.add_argument("--bgm", default=None)
    ap.add_argument("--music-volume", type=float, default=0.12)
    ap.add_argument("--hold", type=float, default=5.0,
                    help="seconds a slide without audio stays on screen (default 5)")
    ap.add_argument("--out", default=None, help="output filename inside the work dir")
    ap.add_argument("--dry-run", action="store_true", help="stop before TTS/video")
    ap.add_argument("--allow-silent", action="store_true",
                    help="build even when narration.json is missing (silent slideshow)")
    ap.add_argument("--strict", action="store_true",
                    help="stop at the first page edge-tts refuses (default: keep the other "
                         "pages, render the video, and report ok:false)")
    args = ap.parse_args(normalize_argv(sys.argv[1:] if argv is None else argv))

    try:
        return {"extract": cmd_extract, "build": cmd_build, "all": cmd_all}[args.phase](args)
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=1))
        return 1


if __name__ == "__main__":
    sys.exit(main())
